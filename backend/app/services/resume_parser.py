"""Extracts plain text from uploaded resume files. Supports PDF (via PyMuPDF),
DOCX (via python-docx) and TXT. Raises ValueError with a user-facing message
for unsupported or unreadable files.
"""
import io
import logging
import re
import zipfile

import fitz  # PyMuPDF
from docx import Document

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt"}


def extract_text(file_bytes: bytes, filename: str) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "doc":
        raise ValueError("Old .doc files aren't supported. Save it as .docx or PDF and upload again.")
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type{f' .{ext}' if ext else ''}. Upload a PDF, DOCX, or TXT file.")
    try:
        if ext == "pdf":
            text = _extract_pdf(file_bytes)
        elif ext == "docx":
            text = _extract_docx(file_bytes)
        else:
            text = _decode_txt(file_bytes)
    except ValueError:
        raise
    except Exception as exc:  # corrupt/encrypted/mislabelled files
        logger.info("Resume parse failed for .%s: %s", ext, type(exc).__name__)
        raise ValueError(f"We couldn't read that {ext.upper()} file. It may be corrupted or password-protected.") from exc
    return _normalize(text)


def _extract_pdf(file_bytes: bytes) -> str:
    if not file_bytes.startswith(b"%PDF"):
        raise ValueError("That file doesn't look like a valid PDF.")
    text_parts = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        if doc.needs_pass:
            raise ValueError("That PDF is password-protected. Remove the password and upload again.")
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def _extract_docx(file_bytes: bytes) -> str:
    if not zipfile.is_zipfile(io.BytesIO(file_bytes)):
        raise ValueError("That file doesn't look like a valid DOCX document.")
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tables (many resumes put skills/experience in table layouts)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs)


def _decode_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("latin-1")


def _normalize(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
